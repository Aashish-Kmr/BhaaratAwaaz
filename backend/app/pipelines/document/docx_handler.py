from __future__ import annotations

from typing import Callable

from docx import Document

from app.core.translation import translator


class DocxHandler:
    def translate(
        self,
        input_file,
        output_file,
        source_language: str,
        target_language: str,
        on_progress: Callable[[float], None] | None = None,
        should_stop: Callable[[], bool] | None = None,
    ):
        document = Document(input_file)

        paragraphs = list(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)

        total = max(1, len(paragraphs))

        for index, paragraph in enumerate(paragraphs):
            if should_stop and should_stop():
                break

            self._translate_paragraph(paragraph, source_language, target_language)

            if on_progress:
                on_progress((index + 1) / total)

        document.save(output_file)

    def _translate_paragraph(self, paragraph, source_language, target_language):
        original_text = paragraph.text

        if not original_text.strip():
            return

        translated = translator.translate(original_text, source_language, target_language)

        if not translated:
            return

        self._replace_text_preserve_formatting(paragraph, translated)

    def _replace_text_preserve_formatting(self, paragraph, translated_text):
        runs = paragraph.runs

        if not runs:
            paragraph.add_run(translated_text)
            return

        formatting = [
            {
                "bold": run.bold,
                "italic": run.italic,
                "underline": run.underline,
                "font_name": run.font.name,
                "font_size": run.font.size,
                "font_color": run.font.color.rgb if run.font.color is not None else None,
                "style": run.style,
            }
            for run in runs
        ]

        original_lengths = [len(run.text) for run in runs]
        total_original_length = sum(original_lengths)

        if len(runs) == 1:
            run = runs[0]
            run.text = translated_text
            self._apply_formatting(run, formatting[0])
            return

        for run in runs:
            run.text = ""

        translated_length = len(translated_text)

        if total_original_length == 0:
            runs[0].text = translated_text
            self._apply_formatting(runs[0], formatting[0])
            return

        positions = []
        current_position = 0

        for index, original_length in enumerate(original_lengths):
            if index == len(runs) - 1:
                end_position = translated_length
            else:
                ratio = original_length / total_original_length
                chunk_length = round(translated_length * ratio)
                end_position = current_position + chunk_length

            positions.append((current_position, end_position))
            current_position = end_position

        for index, run in enumerate(runs):
            start, end = positions[index]
            run.text = translated_text[start:end]
            self._apply_formatting(run, formatting[index])

    def _apply_formatting(self, run, formatting):
        run.bold = formatting["bold"]
        run.italic = formatting["italic"]
        run.underline = formatting["underline"]
        run.font.name = formatting["font_name"]
        run.font.size = formatting["font_size"]

        if formatting["font_color"] is not None:
            run.font.color.rgb = formatting["font_color"]

        if formatting["style"] is not None:
            run.style = formatting["style"]


docx_handler = DocxHandler()
