from __future__ import annotations

from typing import Callable

from openpyxl import load_workbook

from app.core.translation import translator


class XlsxHandler:
    def translate(
        self,
        input_file,
        output_file,
        source_language: str,
        target_language: str,
        on_progress: Callable[[float], None] | None = None,
        should_stop: Callable[[], bool] | None = None,
    ):
        workbook = load_workbook(input_file, data_only=False)

        cells = [
            cell
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
        ]
        total = max(1, len(cells))

        for index, cell in enumerate(cells):
            if should_stop and should_stop():
                break

            value = cell.value

            if isinstance(value, str) and not value.startswith("="):
                text = value.strip()
                if text:
                    cell.value = translator.translate(text, source_language, target_language)

            if on_progress:
                on_progress((index + 1) / total)

        workbook.save(output_file)


xlsx_handler = XlsxHandler()
