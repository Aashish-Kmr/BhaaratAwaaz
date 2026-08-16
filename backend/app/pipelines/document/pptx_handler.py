from __future__ import annotations

from typing import Callable

from pptx import Presentation

from app.core.translation import translator

GROUP_SHAPE_TYPE = 6


class PptxHandler:
    def translate(
        self,
        input_file,
        output_file,
        source_language: str,
        target_language: str,
        on_progress: Callable[[float], None] | None = None,
        should_stop: Callable[[], bool] | None = None,
    ):
        presentation = Presentation(input_file)
        slides = list(presentation.slides)
        total = max(1, len(slides))

        for index, slide in enumerate(slides):
            if should_stop and should_stop():
                break

            for shape in slide.shapes:
                self._process_shape(shape, source_language, target_language)

            if on_progress:
                on_progress((index + 1) / total)

        presentation.save(output_file)

    def _process_shape(self, shape, source_language, target_language):
        if hasattr(shape, "text_frame") and shape.has_text_frame:
            self._translate_text_frame(shape.text_frame, source_language, target_language)

        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    self._translate_text_frame(cell.text_frame, source_language, target_language)

        if shape.shape_type == GROUP_SHAPE_TYPE:
            for child_shape in shape.shapes:
                self._process_shape(child_shape, source_language, target_language)

    def _translate_text_frame(self, text_frame, source_language, target_language):
        for paragraph in text_frame.paragraphs:
            original_text = paragraph.text

            if not original_text.strip():
                continue

            translated = translator.translate(original_text, source_language, target_language)

            if not translated:
                continue

            self._replace_paragraph_text(paragraph, translated)

    def _replace_paragraph_text(self, paragraph, translated_text):
        runs = paragraph.runs

        if not runs:
            run = paragraph.add_run()
            run.text = translated_text
            return

        if len(runs) == 1:
            runs[0].text = translated_text
            return

        formatting = [
            {
                "bold": run.font.bold,
                "italic": run.font.italic,
                "underline": run.font.underline,
                "font_size": run.font.size,
                "font_name": run.font.name,
            }
            for run in runs
        ]

        original_lengths = [len(run.text) for run in runs]
        total_length = sum(original_lengths)

        for run in runs:
            run.text = ""

        if total_length == 0:
            runs[0].text = translated_text
            return

        translated_length = len(translated_text)
        current_position = 0

        for index, run in enumerate(runs):
            if index == len(runs) - 1:
                end_position = translated_length
            else:
                ratio = original_lengths[index] / total_length
                chunk_length = round(translated_length * ratio)
                end_position = current_position + chunk_length

            run.text = translated_text[current_position:end_position]

            fmt = formatting[index]
            run.font.bold = fmt["bold"]
            run.font.italic = fmt["italic"]
            run.font.underline = fmt["underline"]
            run.font.size = fmt["font_size"]
            run.font.name = fmt["font_name"]

            current_position = end_position


pptx_handler = PptxHandler()
