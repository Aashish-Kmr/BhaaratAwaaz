from pptx import Presentation
from services.translation_service import translation_service

BATCH_SIZE = 16


class PptxHandler:
    def translate_pptx(self, input_path, output_path, target_language="hi", source_language=None):
        prs = Presentation(input_path)

        shapes_with_text = []

        # Collect text paragraphs from text frames and tables
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            shapes_with_text.append(paragraph)
                elif shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                shapes_with_text.append(cell)

        if not shapes_with_text:
            prs.save(output_path)
            return

        all_texts = [item.text for item in shapes_with_text]
        translated_texts = []

        # Batch translate across chunks
        for i in range(0, len(all_texts), BATCH_SIZE):
            chunk = all_texts[i:i + BATCH_SIZE]
            translated_texts.extend(
                translation_service.translate_texts(
                    texts=chunk,
                    target_language=target_language,
                    source_language=source_language
                )
            )

        for item, trans_text in zip(shapes_with_text, translated_texts):
            item.text = trans_text

        prs.save(output_path)


pptx_handler = PptxHandler()