from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table

from services.translation_service import translation_service


class DocxHandler:

    def translate(
        self,
        input_file,
        output_file,
        target_language,
    ):
        document = Document(input_file)

        # ---------------------------------------------------------
        # NORMAL DOCUMENT PARAGRAPHS
        # ---------------------------------------------------------
        for paragraph in document.paragraphs:
            self._translate_paragraph(
                paragraph,
                target_language,
            )

        # ---------------------------------------------------------
        # TABLES
        # ---------------------------------------------------------
        for table in document.tables:
            self._translate_table(
                table,
                target_language,
            )

        # ---------------------------------------------------------
        # SAVE
        # ---------------------------------------------------------
        document.save(output_file)

    # =============================================================
    # PARAGRAPH
    # =============================================================

    def _translate_paragraph(
        self,
        paragraph,
        target_language,
    ):
        original_text = paragraph.text

        if not original_text.strip():
            return

        translated = translation_service.translate_text(
            original_text,
            target_language,
        )

        if not translated:
            return

        self._replace_text_preserve_formatting(
            paragraph,
            translated,
        )

    # =============================================================
    # TABLE
    # =============================================================

    def _translate_table(
        self,
        table,
        target_language,
    ):
        for row in table.rows:

            for cell in row.cells:

                for paragraph in cell.paragraphs:

                    self._translate_paragraph(
                        paragraph,
                        target_language,
                    )

    # =============================================================
    # REPLACE TEXT WHILE PRESERVING RUN FORMATTING
    # =============================================================

    def _replace_text_preserve_formatting(
        self,
        paragraph,
        translated_text,
    ):
        runs = paragraph.runs

        # ---------------------------------------------------------
        # No runs
        # ---------------------------------------------------------
        if not runs:
            paragraph.add_run(translated_text)
            return

        # ---------------------------------------------------------
        # Save formatting of every run
        # ---------------------------------------------------------
        formatting = []

        for run in runs:

            formatting.append(
                {
                    "bold": run.bold,
                    "italic": run.italic,
                    "underline": run.underline,

                    "font_name": run.font.name,
                    "font_size": run.font.size,

                    "font_color": (
                        run.font.color.rgb
                        if run.font.color is not None
                        else None
                    ),

                    "style": run.style,
                }
            )

        # ---------------------------------------------------------
        # Determine original text length
        # ---------------------------------------------------------
        original_lengths = [
            len(run.text)
            for run in runs
        ]

        total_original_length = sum(original_lengths)

        # ---------------------------------------------------------
        # If there is only one run, this is easy.
        # ---------------------------------------------------------
        if len(runs) == 1:

            run = runs[0]

            run.text = translated_text

            self._apply_formatting(
                run,
                formatting[0],
            )

            return

        # ---------------------------------------------------------
        # Clear existing text
        # ---------------------------------------------------------
        for run in runs:
            run.text = ""

        # ---------------------------------------------------------
        # Distribute translated text among original runs.
        #
        # We use the original run length as a ratio.
        # ---------------------------------------------------------
        translated_length = len(translated_text)

        if total_original_length == 0:

            runs[0].text = translated_text

            self._apply_formatting(
                runs[0],
                formatting[0],
            )

            return

        positions = []

        current_position = 0

        for index, original_length in enumerate(
            original_lengths
        ):

            if index == len(runs) - 1:

                end_position = translated_length

            else:

                ratio = (
                    original_length
                    / total_original_length
                )

                chunk_length = round(
                    translated_length * ratio
                )

                end_position = (
                    current_position
                    + chunk_length
                )

            positions.append(
                (
                    current_position,
                    end_position,
                )
            )

            current_position = end_position

        # ---------------------------------------------------------
        # Put translated text back into runs
        # ---------------------------------------------------------
        for index, run in enumerate(runs):

            start, end = positions[index]

            run.text = translated_text[
                start:end
            ]

            self._apply_formatting(
                run,
                formatting[index],
            )

    # =============================================================
    # APPLY RUN FORMATTING
    # =============================================================

    def _apply_formatting(
        self,
        run,
        formatting,
    ):
        run.bold = formatting["bold"]
        run.italic = formatting["italic"]
        run.underline = formatting["underline"]

        run.font.name = formatting["font_name"]
        run.font.size = formatting["font_size"]

        if formatting["font_color"] is not None:

            run.font.color.rgb = (
                formatting["font_color"]
            )

        if formatting["style"] is not None:

            run.style = formatting["style"]


docx_handler = DocxHandler()